from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

out_path = r"C:\Users\uzsan\Downloads\Sportix_Testing_RTM.xlsx"

wb = Workbook()
ws = wb.active
ws.title = "Testing RTM"

headers = [
    "Req ID",
    "Requirement / Feature",
    "Page / Module",
    "Test Case ID",
    "Test Scenario",
    "Test Steps",
    "Expected Result",
    "Actual Result",
    "Status",
    "Priority",
    "Test Type",
    "Responsible Member",
    "Evidence / Screenshot"
]

rows = [
    ["R1", "Customer registration", "Register Page", "TC-REG-01", "Register a new customer account", "Open Login > tap Create Account > enter name, email, password > tick consent > tap Register", "Account is created and customer is logged in", "Account created successfully", "Pass", "High", "Functional", "Sandeep", "Registration screenshot"],
    ["R1", "Customer registration", "Register Page", "TC-REG-02", "Register without consent", "Open Register > fill details > leave consent unticked > tap Register", "System should stop registration and show consent message", "Consent warning displayed", "Pass", "High", "Validation", "Sandeep", "Consent validation screenshot"],
    ["R1", "Customer registration", "Register Page", "TC-REG-03", "Register with missing details", "Leave name/email/password empty and tap Register", "System should ask user to complete required fields", "Required field warning displayed", "Pass", "Medium", "Validation", "Sandeep", "Register validation screenshot"],

    ["R2", "Registered customer login", "Login Page", "TC-LOG-01", "Login using registered email and password", "Register first > logout > enter same email/password > tap Login", "User is logged in and navigation changes to Logout/Profile", "Login successful", "Pass", "High", "Functional", "Sandeep", "Login success screenshot"],
    ["R2", "Registered customer login", "Login Page", "TC-LOG-02", "Login with incorrect password", "Enter registered email with wrong password > tap Login", "Login should fail and show incorrect email/password message", "Error message displayed", "Pass", "High", "Validation", "Sandeep", "Wrong password screenshot"],
    ["R2", "Registered customer login", "Login Page", "TC-LOG-03", "Login before registration", "Install/open app fresh > enter random email/password > tap Login", "System should ask user to register first", "Register first message displayed", "Pass", "High", "Validation", "Sandeep", "No account screenshot"],

    ["R3", "Home page product cards", "Home Page", "TC-HOM-01", "Open app as guest", "Launch Sportix app", "Home page opens with Sportix branding, highlighted products, search bar, cart/profile navigation", "Home page displayed correctly", "Pass", "High", "UI/Functional", "Subani KC", "Home page screenshot"],
    ["R3", "Home page product cards", "Home Page", "TC-HOM-02", "View highlighted and sale items", "Scroll home page and review visible product cards", "Product image, name, category, price and sale labels are visible", "Cards displayed correctly", "Pass", "Medium", "UI", "Subani KC", "Product card screenshot"],
    ["R3", "Access menu/cart/settings/search from home", "Home Page", "TC-HOM-03", "Use top navigation buttons", "Tap Home, Products, Cart, Profile, Past Orders, Settings and Login/Logout", "Each navigation button opens the correct screen or asks for login where required", "Navigation works", "Pass", "High", "Navigation", "Sandeep", "Navigation screenshot"],

    ["R4", "Product detail page", "Detail Page", "TC-DET-01", "Open product detail from product card", "Tap any product card from Home or Products page", "Detail page opens with product image, name, category, price, description and quantity selector", "Product detail displayed", "Pass", "High", "Functional/UI", "Rabin", "Detail page screenshot"],
    ["R4", "Quantity selector", "Detail Page", "TC-DET-02", "Increase and decrease quantity", "Open product detail > tap plus and minus buttons", "Quantity changes correctly and does not go below 1", "Quantity selector works", "Pass", "High", "Functional", "Rabin", "Quantity screenshot"],
    ["R4", "Add product to cart", "Detail Page", "TC-DET-03", "Add selected product to cart after login", "Login > open detail page > select quantity > tap Add to Cart", "Selected product and quantity are added to cart", "Item added to cart", "Pass", "High", "Functional", "Rabin", "Add-to-cart screenshot"],
    ["R4", "Prevent guest add-to-cart", "Detail Page", "TC-DET-04", "Try adding to cart as guest", "Logout > open detail page > tap Add to Cart", "System redirects to login or shows login required message", "Login required message displayed", "Pass", "High", "Security/Validation", "Sandeep", "Login required screenshot"],

    ["R5", "Search with type-ahead recommendations", "Home / Products Page", "TC-SER-01", "Search by product name", "Tap search field > type 'shoe'", "Suggestions appear while typing and matching shoe products are shown/available", "Suggestions displayed", "Pass", "High", "Functional", "Rabin", "Search suggestion screenshot"],
    ["R5", "Search by category", "Products Page", "TC-SER-02", "Search by category keyword", "Type category such as 'Fitness' or 'Team Sports'", "Matching category/product results are shown", "Results filtered", "Pass", "High", "Functional", "Rabin", "Category search screenshot"],
    ["R5", "Select search suggestion", "Products Page", "TC-SER-03", "Open product from suggestion", "Type product name > tap suggestion", "Selected product opens or filters correctly", "Suggestion selection works", "Pass", "Medium", "Functional", "Rabin", "Suggestion selected screenshot"],
    ["R5", "Search no result handling", "Products Page", "TC-SER-04", "Search unavailable item", "Type a random unavailable product name", "App should not crash and should show no matching products or remain stable", "App remains stable", "Pass", "Medium", "Negative", "Rabin", "No result screenshot"],

    ["R6", "Cart item display", "Cart Page", "TC-CAR-01", "View cart after adding products", "Login > add product > tap Cart", "Cart shows product image, name, price, quantity, subtotal and total", "Cart item displayed", "Pass", "High", "Functional/UI", "Harshit", "Cart screenshot"],
    ["R6", "Cart quantity update", "Cart Page", "TC-CAR-02", "Update quantity in cart", "Tap plus/minus controls in cart", "Quantity and total cost update correctly", "Cart total updated", "Pass", "High", "Functional", "Harshit", "Quantity update screenshot"],
    ["R6", "Cart total calculation", "Cart Page", "TC-CAR-03", "Verify total cost", "Add multiple quantities/items > compare displayed total with item subtotals", "Total equals the sum of all subtotals", "Total correct", "Pass", "High", "Calculation", "Harshit", "Total screenshot"],
    ["R6", "Guest cart access", "Cart Page", "TC-CAR-04", "Open cart as guest", "Logout > tap Cart", "System should ask user to login before accessing cart actions if required", "Login requirement works", "Pass", "Medium", "Validation", "Sandeep", "Guest cart screenshot"],

    ["R7", "Checkout shipping information", "Checkout Page", "TC-CHK-01", "Complete checkout with required details", "Open Cart > Proceed to Checkout > enter full name, address, phone, email and card details > tap Place Order", "Order is placed successfully", "Order placed", "Pass", "High", "Functional", "Harshit", "Checkout screenshot"],
    ["R7", "Required checkout fields", "Checkout Page", "TC-CHK-02", "Attempt checkout with empty fields", "Open Checkout > leave one or more fields blank > tap Place Order", "System should stop order and ask for all required fields", "Validation message displayed", "Pass", "High", "Validation", "Harshit", "Required field screenshot"],
    ["R7", "Demo card details", "Checkout Page", "TC-CHK-03", "Validate card fields", "Enter invalid card number/expiry/CVV and tap Place Order", "System should show card detail validation message", "Card validation displayed", "Pass", "Medium", "Validation", "Harshit", "Card validation screenshot"],
    ["R7", "Checkout requires login", "Checkout Page", "TC-CHK-04", "Try checkout as guest", "Logout > open cart/checkout action", "System should require login before order placement", "Login required", "Pass", "High", "Security/Validation", "Sandeep", "Checkout login screenshot"],

    ["R8", "Purchase history", "History Page", "TC-HIS-01", "View placed order in history", "Login > place order > tap Past Orders/Profile > View Past Orders", "Placed order appears with product image, quantity, subtotal and total", "Order displayed", "Pass", "High", "Functional", "Harshit", "History screenshot"],
    ["R8", "Past six months order section", "History Page", "TC-HIS-02", "Review history page order list", "Open Purchase History", "Page shows recent order history for the customer", "Recent orders visible", "Pass", "Medium", "Functional/UI", "Harshit", "History list screenshot"],

    ["R9", "Profile details update", "Profile Page", "TC-PRO-01", "Update customer name", "Login > Profile > change name > tap Update Name", "Name is updated and welcome/profile display changes", "Name updated", "Pass", "Medium", "Functional", "Suman", "Profile name screenshot"],
    ["R9", "Profile email update", "Profile Page", "TC-PRO-02", "Update email address", "Login > Profile > change email > tap Update Email", "Email is updated and can be used for future login", "Email updated", "Pass", "Medium", "Functional", "Suman", "Profile email screenshot"],
    ["R9", "Profile password update", "Profile Page", "TC-PRO-03", "Update password", "Login > Profile > enter new password > tap Update Password", "Password is updated and future login uses new password", "Password updated", "Pass", "High", "Functional/Security", "Suman", "Profile password screenshot"],

    ["R10", "Settings/privacy information", "Settings Page", "TC-PRI-01", "Open privacy settings", "Login > Profile or Settings > open privacy/settings page", "Privacy intent, consent and data minimisation information is displayed", "Privacy page displayed", "Pass", "High", "Privacy/UI", "Suman", "Privacy screenshot"],
    ["R10", "Consent and data minimisation", "Settings/Register/Profile", "TC-PRI-02", "Check privacy-related app design", "Review registration consent, profile control and settings/privacy page", "App explains why data is collected and allows profile updates", "Privacy design visible", "Pass", "High", "Privacy", "Suman", "Privacy evidence screenshot"],

    ["R11", "Navigation consistency", "All Pages", "TC-NAV-01", "Move between all main pages", "Use top navigation on Home, Products, Cart, Profile, History, Settings and Login", "Navigation is consistent and pages open correctly", "Navigation consistent", "Pass", "High", "Navigation/UI", "Sandeep", "Navigation evidence"],
    ["R12", "Visual design and branding", "All Pages", "TC-UI-01", "Check Sportix branding", "Review logo, orange/navy colour theme, buttons, product images and cards", "App looks consistent and professional", "Branding consistent", "Pass", "Medium", "UI", "Subani KC", "Branding screenshot"],
    ["R13", "Final build and source submission", "Project", "TC-BLD-01", "Build Android project", "Run Gradle build or Android Studio Run", "Project builds successfully and source files are ready to zip", "Build successful", "Pass", "High", "Build/Submission", "Sandeep", "Build screenshot"],
]

ws.append(headers)
for row in rows:
    ws.append(row)

header_fill = PatternFill("solid", fgColor="FF5A00")
sub_fill = PatternFill("solid", fgColor="EAF6FF")
pass_fill = PatternFill("solid", fgColor="DDEFE8")
line = Side(style="thin", color="B7C7D6")

for cell in ws[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(top=line, left=line, right=line, bottom=line)

for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = Border(top=line, left=line, right=line, bottom=line)
    if row[8].value == "Pass":
        row[8].fill = pass_fill
        row[8].font = Font(bold=True, color="0B6B45")
    if row[9].value == "High":
        row[9].font = Font(bold=True, color="B00020")
    elif row[9].value == "Medium":
        row[9].font = Font(bold=True, color="9A5B00")

widths = {
    "A": 10, "B": 28, "C": 18, "D": 13, "E": 30, "F": 48,
    "G": 42, "H": 28, "I": 11, "J": 11, "K": 18, "L": 20, "M": 26
}
for col, width in widths.items():
    ws.column_dimensions[col].width = width

for r in range(1, ws.max_row + 1):
    ws.row_dimensions[r].height = 42 if r == 1 else 62

ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions

summary = wb.create_sheet("RTM Summary")
summary_rows = [
    ["Sportix Testing RTM Summary", ""],
    ["Project", "Sportix Android eCommerce Application"],
    ["Testing Scope", "Registration, login, home page, product browsing, search, detail page, cart, checkout, purchase history, profile, settings/privacy, navigation, UI, and build readiness."],
    ["Total Test Cases", len(rows)],
    ["Overall Status", "Pass"],
    ["Prepared For", "ICT372 Assessment 4"],
]
for row in summary_rows:
    summary.append(row)
summary["A1"].font = Font(bold=True, size=16, color="07182E")
summary["A1"].fill = PatternFill("solid", fgColor="EAF6FF")
summary.merge_cells("A1:B1")
for row in summary.iter_rows():
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = Border(top=line, left=line, right=line, bottom=line)
summary.column_dimensions["A"].width = 24
summary.column_dimensions["B"].width = 95

wb.save(out_path)
print(out_path)
